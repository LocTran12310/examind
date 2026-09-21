"""Bulk account import from CSV/XLSX (AC-16, AC-17, A-08)."""
import csv
from dataclasses import dataclass, field
import io

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.errors import AppError, validation
from app.core.passwords import temp_password
from app.core.security import hash_password
from app.deps import OrgScope
from app.models import ClassMember, User
from app.services import audit, classes as class_service
from app.services.users import ORG_ROLES, USERNAME_RE, can_manage, unique_username

MAX_ROWS = 2000
COLUMNS = {"full_name", "username", "role", "class"}
ALIASES = {"họ tên": "full_name", "ho ten": "full_name", "họ và tên": "full_name", "ten dang nhap": "username",
           "tên đăng nhập": "username", "vai trò": "role", "lớp": "class", "lop": "class"}
ROLE_ALIASES = {"hs": "student", "học sinh": "student", "gv": "teacher", "giáo viên": "teacher", "admin": "org_admin"}


@dataclass
class Row:
    row: int
    full_name: str = ""
    username: str = ""
    role: str = "student"
    klass: str = ""
    errors: list[str] = field(default_factory=list)
    generated: bool = False

    def as_dict(self):
        return {"row": self.row, "full_name": self.full_name, "username": self.username, "role": self.role,
                "class": self.klass, "errors": self.errors, "generated_username": self.generated}


def _header_key(h: str) -> str:
    h = (h or "").strip().lower().lstrip("﻿")
    return ALIASES.get(h, h)


def parse_file(filename: str, data: bytes) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
    elif name.endswith(".csv") or name.endswith(".txt"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise validation("File phải mã hóa UTF-8", "file")
        first = text.splitlines()[0] if text.strip() else ""
        delimiter = max(",;\t", key=first.count) if first else ","
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    else:
        raise validation("Chỉ hỗ trợ file .csv hoặc .xlsx", "file")
    rows = [r for r in rows if any((c or "").strip() for c in r)]
    if not rows:
        raise validation("File rỗng", "file")
    header = [_header_key(h) for h in rows[0]]
    if "full_name" not in header:
        raise validation("Thiếu cột full_name (họ tên)", "file")
    if len(rows) - 1 > MAX_ROWS:
        raise validation(f"Tối đa {MAX_ROWS} dòng mỗi lần", "file")
    out = []
    for r in rows[1:]:
        out.append({k: (r[i].strip() if i < len(r) and r[i] else "") for i, k in enumerate(header) if k in COLUMNS})
    return out


def validate_rows(db: Session, scope: OrgScope, raw_rows: list[dict]) -> list[Row]:
    rows: list[Row] = []
    for i, r in enumerate(raw_rows, start=2):  # row 1 is the header
        role = (r.get("role") or "student").strip().lower()
        role = ROLE_ALIASES.get(role, role)
        rows.append(Row(row=r.get("row", i), full_name=(r.get("full_name") or "").strip(),
                        username=(r.get("username") or "").strip().lower(), role=role, klass=(r.get("class") or "").strip()))
    given = [r.username for r in rows if r.username]
    existing = set(u.lower() for u in db.scalars(select(User.username).where(User.organization_id == scope.org_id, User.username.in_(given)))) if given else set()
    seen: set[str] = set()
    for r in rows:
        if not r.full_name:
            r.errors.append("Thiếu họ tên")
        if r.role not in ORG_ROLES:
            r.errors.append(f"Vai trò không hợp lệ: {r.role}")
        elif not can_manage(scope, r.role):
            r.errors.append("Bạn không có quyền tạo vai trò này")
        if r.username:
            if not USERNAME_RE.match(r.username):
                r.errors.append("Tên đăng nhập không hợp lệ")
            elif r.username in existing:
                r.errors.append("Tên đăng nhập đã tồn tại")
            elif r.username in seen:
                r.errors.append("Tên đăng nhập bị trùng trong file")
            seen.add(r.username)
    for r in rows:
        if not r.username and r.full_name:
            r.username = unique_username(db, scope.org_id, r.full_name, taken=seen)
            r.generated = True
            seen.add(r.username)
    return rows


def commit(db: Session, scope: OrgScope, raw_rows: list[dict]) -> list[dict]:
    rows = validate_rows(db, scope, raw_rows)
    bad = [r.as_dict() for r in rows if r.errors]
    if bad:
        raise AppError("import_invalid", f"{len(bad)} dòng có lỗi, chưa tạo tài khoản nào", 422, {"rows": bad})
    created, class_cache = [], {}
    for r in rows:
        password = temp_password()
        user = User(organization_id=scope.org_id, username=r.username, full_name=r.full_name, role=r.role,
                    password_hash=hash_password(password), must_change_password=True)
        db.add(user)
        db.flush()
        if r.klass:
            if r.klass not in class_cache:
                class_cache[r.klass] = class_service.find_or_create(db, scope, r.klass)
            db.add(ClassMember(class_id=class_cache[r.klass].id, user_id=user.id))
        created.append({"user_id": user.id, "username": r.username, "full_name": r.full_name, "role": r.role,
                        "class": r.klass, "temp_password": password})
    audit.record(db, scope.user, scope.org_id, "user.import", "user", None, count=len(created))
    return created
