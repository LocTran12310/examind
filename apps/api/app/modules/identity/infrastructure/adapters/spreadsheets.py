"""SpreadsheetReader over openpyxl (.xlsx) and the csv module (.csv/.txt, UTF-8, delimiter sniffed from the header)."""
import csv
import io

from openpyxl import load_workbook

from app.shared.domain.errors import Invalid


class FileSpreadsheetReader:
    def table(self, filename: str, data: bytes) -> list[list[str]]:
        name = (filename or "").lower()
        if name.endswith(".xlsx"):
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            return [["" if c is None else str(c) for c in r] for r in wb.active.iter_rows(values_only=True)]
        if name.endswith(".csv") or name.endswith(".txt"):
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                raise Invalid("File phải mã hóa UTF-8", "file")
            first = text.splitlines()[0] if text.strip() else ""
            delimiter = max(",;\t", key=first.count) if first else ","
            return list(csv.reader(io.StringIO(text), delimiter=delimiter))
        raise Invalid("Chỉ hỗ trợ file .csv hoặc .xlsx", "file")
